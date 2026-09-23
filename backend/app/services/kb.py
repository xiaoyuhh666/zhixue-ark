"""知识库服务（里程碑 4）：文档解析 -> 分块 -> API 嵌入 -> 向量表检索。

设计要点：
- 嵌入走智谱 embedding-3（OpenAI 兼容 /embeddings，复用 GLM_API_KEY），不再本地加载模型
  —— 免去 torch/sentence-transformers 依赖与 ~2GB 内存，Render 免费层 512MB 稳定运行
- 分批请求（智谱单请求上限 64 条）+ L2 归一化，返回顺序与输入一致
- 向量存储 2026-09-23 起用数据库表 kb_chunks（ChromaDB 退役）：
  · 配置 Turso 时走远端 libSQL，embedding 列为 F32_BLOB 向量类型，
    检索用服务端 vector_distance_cos 计算，数据重启不丢
  · 未配置 Turso 时同一张表落在本地 SQLite（F32_BLOB 退化为普通 BLOB），
    相似度在 Python 侧用点积计算（向量已归一化，点积即余弦），本地开发零依赖
- 引用溯源：每个向量块带 doc_id/seq，检索命中即可回溯文档名与原文片段，按 seq 还原全文
"""
import json

from sqlalchemy import text

from ..config import get_settings
from ..db import IS_LOCAL_SQLITE, engine
from ..models import KnowledgeDoc

# 解析后拼接段落时的分隔符
_PARA_SEP = "\n"

# 惰性单例的嵌入客户端（进程内复用 HTTP 连接池，避免并发重复创建）
_client = None

# kb_chunks 表是否已确认存在（进程内建一次即可）
_TABLE_READY = False


def _get_client():
    """懒加载智谱嵌入客户端（OpenAI 兼容协议）。"""
    global _client
    if _client is None:
        from openai import OpenAI

        s = get_settings()
        if not s.GLM_API_KEY:
            raise RuntimeError("未配置智谱 API Key（GLM_API_KEY），知识库嵌入不可用")
        _client = OpenAI(api_key=s.GLM_API_KEY, base_url=s.GLM_BASE_URL)
    return _client


# ---------- 向量块表（ChromaDB 替代） ----------

def _ensure_table() -> None:
    """幂等建 kb_chunks 向量块表（进程内确认一次）。

    embedding 列声明为 libSQL 向量类型 F32_BLOB：远端 Turso 上是真正的向量列，
    本地 SQLite 按亲和性规则退化为普通 BLOB，同一套 SQL 两种库通用。
    """
    global _TABLE_READY
    if _TABLE_READY:
        return
    with engine.begin() as conn:
        conn.execute(text(
            """
            CREATE TABLE IF NOT EXISTS kb_chunks (
                id        INTEGER PRIMARY KEY AUTOINCREMENT,
                doc_id    INTEGER NOT NULL,
                seq       INTEGER NOT NULL,
                filename  TEXT    NOT NULL DEFAULT '',
                category  TEXT    NOT NULL DEFAULT 'general',
                content   TEXT    NOT NULL,
                embedding F32_BLOB NOT NULL
            )
            """
        ))
    _TABLE_READY = True


def _pack_vec(vec: list[float]) -> bytes:
    """float32 小端打包（libSQL F32_BLOB 的存储格式，本地 SQLite 同格式兼容）。"""
    import struct

    return struct.pack(f"<{len(vec)}f", *vec)


def _unpack_vec(blob) -> list[float]:
    """把 float32 BLOB 解包回 Python 列表（仅本地 SQLite 检索路径使用）。"""
    import struct

    blob = bytes(blob)
    n = len(blob) // 4
    return list(struct.unpack(f"<{n}f", blob))


def _doc_filter_sql(category: str | None, doc_ids: list[int] | None) -> tuple[str, dict]:
    """构造过滤条件与绑定参数（分类 + 文档范围），两种库共用。"""
    conds, params = ["1=1"], {}
    if category:
        conds.append("category = :category")
        params["category"] = category
    if doc_ids:  # 文档范围过滤：用户在对话页知识库菜单勾选的资料
        keys = []
        for i, d in enumerate(doc_ids):
            params[f"d{i}"] = int(d)
            keys.append(f":d{i}")
        conds.append(f"doc_id IN ({', '.join(keys)})")
    return " AND ".join(conds), params


# ---------- 文档解析 ----------

def _parse_pdf(raw: bytes) -> str:
    from io import BytesIO

    from pypdf import PdfReader

    reader = PdfReader(BytesIO(raw))
    return _PARA_SEP.join(page.extract_text() or "" for page in reader.pages)


def _parse_docx(raw: bytes) -> str:
    from io import BytesIO

    import docx

    d = docx.Document(BytesIO(raw))
    return _PARA_SEP.join(p.text for p in d.paragraphs if p.text.strip())


def _parse_pptx(raw: bytes) -> str:
    from io import BytesIO

    from pptx import Presentation

    prs = Presentation(BytesIO(raw))
    parts = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    texts.append(t)
        if texts:
            parts.append(_PARA_SEP.join(texts))
    return _PARA_SEP.join(parts)


def parse_file(filename: str, raw: bytes) -> str:
    """按扩展名解析文件为纯文本。"""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "pdf":
        return _parse_pdf(raw)
    if ext == "docx":
        return _parse_docx(raw)
    if ext == "pptx":
        return _parse_pptx(raw)
    if ext in ("md", "txt"):
        return raw.decode("utf-8", errors="replace")
    raise ValueError(f"暂不支持 .{ext} 格式，请上传 PDF/DOCX/PPTX/MD/TXT")


# ---------- 分块 ----------

def split_chunks(text: str, size: int | None = None, overlap: int | None = None) -> list[str]:
    """按段落聚合切块：目标 size 字符，相邻块 overlap 重叠，尽量在句子边界切。

    中文文档按字符数即可；优先按换行分段，超长段落再按句号/问号切。
    """
    s = get_settings()
    size = size or s.KB_CHUNK_SIZE
    overlap = overlap or s.KB_CHUNK_OVERLAP
    text = text.strip()
    if not text:
        return []

    # 段落优先，超长段落降级为句子，最后按字符硬切
    paras = [p.strip() for p in text.split("\n") if p.strip()]
    units: list[str] = []  # 最小不可切分单元
    for p in paras:
        if len(p) <= size:
            units.append(p)
            continue
        import re

        sentences = re.split(r"(?<=[。！？!?；;\n])", p)
        buf = ""
        for sent in sentences:
            if not sent:
                continue
            if len(buf) + len(sent) <= size:
                buf += sent
            else:
                if buf:
                    units.append(buf)
                # 单句仍超长则按字符硬切
                while len(sent) > size:
                    units.append(sent[:size])
                    sent = sent[size - overlap:]
                buf = sent
        if buf:
            units.append(buf)

    # 相邻单元聚合到 size 大小
    chunks: list[str] = []
    buf = ""
    for u in units:
        if len(buf) + len(u) + 1 <= size or not buf:
            buf = f"{buf}{_PARA_SEP}{u}" if buf else u
        else:
            chunks.append(buf)
            # 保留 overlap 尾部，保证跨块语义连续
            tail = buf[-overlap:] if overlap > 0 else ""
            buf = f"{tail}{_PARA_SEP}{u}" if tail else u
    if buf:
        chunks.append(buf)
    return [c.strip() for c in chunks if c.strip()]


# ---------- 入库 / 检索 / 删除 ----------

def embed_texts(texts: list[str]) -> list[list[float]]:
    """批量嵌入：调智谱 embedding API，分批请求，返回 L2 归一化向量列表。

    请求按 KB_EMBED_BATCH 分批；兼容实现可能乱序返回，按 index 回填保证顺序一致。
    L2 归一化与旧本地模型行为对齐，保证 KB_MIN_SCORE 余弦阈值语义不变。
    """
    if not texts:
        return []
    s = get_settings()
    client = _get_client()
    vecs: list[list[float]] = []
    for i in range(0, len(texts), s.KB_EMBED_BATCH):
        batch = texts[i : i + s.KB_EMBED_BATCH]
        kwargs = {"model": s.KB_EMBED_MODEL, "input": batch}
        if s.KB_EMBED_DIMENSIONS > 0:
            kwargs["dimensions"] = s.KB_EMBED_DIMENSIONS
        resp = client.embeddings.create(**kwargs)
        batch_vecs: list[list[float] | None] = [None] * len(batch)
        for item in resp.data:
            batch_vecs[item.index] = item.embedding
        if any(v is None for v in batch_vecs):
            raise RuntimeError("嵌入接口返回数据不完整")
        vecs.extend(batch_vecs)
    return [_l2_normalize(v) for v in vecs]


def _l2_normalize(vec: list[float]) -> list[float]:
    """L2 归一化（纯 Python，避免为此引入 numpy）。"""
    norm = sum(x * x for x in vec) ** 0.5
    if norm == 0:
        return vec
    return [x / norm for x in vec]


def add_document(doc: KnowledgeDoc, raw_text: str) -> int:
    """对文档分块并向量化入库，返回块数。每个块带 doc_id/seq，可整篇删除、按序还原。

    参数名用 raw_text 而非 text，避免遮蔽 sqlalchemy.text。
    """
    chunks = split_chunks(raw_text)
    if not chunks:
        return 0
    vecs = embed_texts(chunks)
    _ensure_table()
    with engine.begin() as conn:
        for i, (c, v) in enumerate(zip(chunks, vecs)):
            conn.execute(
                text(
                    "INSERT INTO kb_chunks (doc_id, seq, filename, category, content, embedding) "
                    "VALUES (:doc_id, :seq, :filename, :category, :content, :emb)"
                ),
                {
                    "doc_id": doc.id,
                    "seq": i,
                    "filename": doc.filename,
                    "category": doc.category or "general",  # 元数据带分类：供定向检索过滤
                    "content": c,
                    "emb": _pack_vec(v),
                },
            )
    return len(chunks)


def delete_document(doc_id: int) -> None:
    """按 doc_id 删除该文档的全部向量块。"""
    _ensure_table()
    with engine.begin() as conn:
        conn.execute(text("DELETE FROM kb_chunks WHERE doc_id = :d"), {"d": doc_id})


def get_document_content(doc_id: int) -> str:
    """按 doc_id 取回全部向量块并按 seq 拼接，还原文档全文（里程碑 7 引用溯源）。"""
    _ensure_table()
    with engine.connect() as conn:
        rows = conn.execute(
            text("SELECT content FROM kb_chunks WHERE doc_id = :d ORDER BY seq"),
            {"d": doc_id},
        ).fetchall()
    return "\n".join(r[0] for r in rows)


def get_chunks(doc_id: int) -> list[dict]:
    """按 doc_id 取回全部分块（按块序），供前端「分块可视化」展示。

    返回 [{index, text, chars}]。
    """
    _ensure_table()
    with engine.connect() as conn:
        rows = conn.execute(
            text("SELECT seq, content FROM kb_chunks WHERE doc_id = :d ORDER BY seq"),
            {"d": doc_id},
        ).fetchall()
    return [{"index": seq, "text": t, "chars": len(t)} for seq, t in rows]


def search(
    query: str, k: int | None = None, category: str | None = None,
    doc_ids: list[int] | None = None,
) -> list[dict]:
    """语义检索：返回 [{doc_id, doc, snippet, score}]，score 为归一化余弦相似度。

    category（里程碑 6）：传入分类时 SQL 过滤，实现定向检索。
    doc_ids（2026-09-20）：传入文档 id 列表时只在这些文档的块中检索（用户在
    对话页知识库菜单里勾选资料）；None/空列表表示不过滤（检索全部）。

    两种检索路径：
    - 远端 libSQL（Turso）：vector_distance_cos 在服务端算距离，SQL 取 Top-K
    - 本地 SQLite：取出候选块，Python 侧点积排序（向量已归一化，点积即余弦）
    """
    s = get_settings()
    _ensure_table()
    k = k or s.KB_TOP_K
    where_sql, params = _doc_filter_sql(category, doc_ids)

    with engine.connect() as conn:
        if not IS_LOCAL_SQLITE:
            # 远端路径：嵌入后传 JSON 给 vector32，服务端计算余弦距离
            qv = embed_texts([query])[0]
            rows = conn.execute(
                text(
                    f"SELECT doc_id, filename, content, "
                    f"vector_distance_cos(embedding, vector32(:qv)) AS dist "
                    f"FROM kb_chunks WHERE {where_sql} ORDER BY dist ASC LIMIT {int(k)}"
                ),
                {**params, "qv": json.dumps(qv)},
            ).fetchall()
            hits = []
            for doc_id, filename, content, dist in rows:
                score = round(1.0 - float(dist), 4)  # 余弦距离 -> 相似度（与旧 ChromaDB 语义一致）
                if score < s.KB_MIN_SCORE:
                    continue
                hits.append({
                    "doc_id": doc_id,
                    "doc": filename or "",
                    "snippet": content[:200],
                    "score": score,
                })
            return hits

        # 本地路径：取出全部候选块，Python 点积排序
        rows = conn.execute(
            text(
                f"SELECT doc_id, filename, content, embedding "
                f"FROM kb_chunks WHERE {where_sql}"
            ),
            params,
        ).fetchall()
    if not rows:
        return []

    qv = embed_texts([query])[0]
    scored = []
    for doc_id, filename, content, emb in rows:
        try:
            vec = _unpack_vec(emb)
        except Exception:  # 维度不符/脏数据跳过，不影响其余块
            continue
        score = sum(a * b for a, b in zip(qv, vec))  # 归一化向量的点积即余弦相似度
        if score < s.KB_MIN_SCORE:
            continue
        scored.append((score, doc_id, filename or "", content))
    scored.sort(key=lambda x: x[0], reverse=True)
    return [
        {"doc_id": d, "doc": f, "snippet": c[:200], "score": round(sc, 4)}
        for sc, d, f, c in scored[:k]
    ]
